"""Build and read the 200 human-typed validation queries."""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from bench.env import QUERY_XLSX, SOURCE_XLSX
from bench.query_list import QUERIES
from bench.text_cleaner import clean_text_readable


def load_titles(path: Path) -> dict[int, str]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["JobTitles"] if "JobTitles" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise SystemExit(f"{path.name} has no header")
        columns = [str(cell).strip() if cell is not None else "" for cell in header]
        id_index = columns.index("id")
        title_index = columns.index("title")
        titles: dict[int, str] = {}
        for values in rows:
            if values is None or values[id_index] is None:
                continue
            raw_id = values[id_index]
            row_id = int(raw_id) if not isinstance(raw_id, str) else int(raw_id.strip())
            title = "" if values[title_index] is None else str(values[title_index]).strip()
            titles[row_id] = title
        return titles
    finally:
        workbook.close()


def write_query_workbook() -> Path:
    if len(QUERIES) != 200:
        raise SystemExit(f"expected 200 queries, found {len(QUERIES)}")
    ids = [row_id for row_id, _, _ in QUERIES]
    if len(ids) != len(set(ids)):
        raise SystemExit("duplicate gold id in the query list")

    titles = load_titles(SOURCE_XLSX)
    missing = [row_id for row_id in ids if row_id not in titles]
    if missing:
        raise SystemExit(f"ids not in the spreadsheet: {missing}")

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "ValidationQueries"
    sheet.append(["query", "gold_id", "gold_title", "change"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row_id, query, change in QUERIES:
        sheet.append([query, row_id, titles[row_id], change])
    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"
    for column, width in {"A": 48, "B": 12, "C": 56, "D": 16}.items():
        sheet.column_dimensions[column].width = width
    QUERY_XLSX.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(QUERY_XLSX)
    return QUERY_XLSX


def read_queries(path: Path = QUERY_XLSX) -> list[dict[str, object]]:
    if not path.is_file():
        raise SystemExit(f"Query file not found: {path.name}. Run: python -m bench queries")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
        columns = {name: index for index, name in enumerate(header)}
        for required in ("query", "gold_id", "gold_title"):
            if required not in columns:
                raise SystemExit(f"{path.name} is missing a {required} column")
        records = []
        for values in rows:
            if values is None or values[columns["query"]] is None:
                continue
            query = str(values[columns["query"]]).strip()
            if not query:
                continue
            raw_id = values[columns["gold_id"]]
            gold_id = int(raw_id) if not isinstance(raw_id, str) else int(str(raw_id).strip())
            gold_title = "" if values[columns["gold_title"]] is None else str(values[columns["gold_title"]]).strip()
            records.append(
                {
                    "gold_id": gold_id,
                    "query": query,
                    "query_normalized": clean_text_readable(query),
                    "gold_title": gold_title,
                }
            )
        return records
    finally:
        workbook.close()
